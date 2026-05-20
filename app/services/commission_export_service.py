"""Exportación Excel/PDF de comisiones (P21-G)."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from sqlalchemy.orm import Session

from app.core.paths import COMMISSION_EXPORTS_DIR
from app.models.commission import PAYMENT_STATUS_CANCELLED
from app.repositories.commission_repository import CommissionRepository

_STATUS_LABELS = {
    "pending": "Pendiente",
    "paid": "Pagada",
    "cancelled": "Cancelada",
}


def _rows_for_export(db: Session, **filters) -> list:
    commissions = CommissionRepository.list_commissions(db, **filters, limit=5000)
    rows = []
    for c in commissions:
        sale = c.sale_capture
        rows.append(
            {
                "folio": sale.folio if sale else "—",
                "cliente": sale.cliente if sale else "—",
                "vendedor": c.seller_name,
                "qna": c.qna or "",
                "semana": c.week or "",
                "venta": c.sale_amount,
                "pct": c.commission_percentage,
                "comision": c.commission_amount,
                "estado": _STATUS_LABELS.get(c.payment_status, c.payment_status),
            }
        )
    return rows


def _summary_by_seller_qna(rows: list[dict]) -> list[dict]:
    agg: dict[tuple[str, str], Decimal] = {}
    for r in rows:
        if r["estado"] == "Cancelada":
            continue
        key = (r["vendedor"], r["qna"] or "—")
        agg[key] = agg.get(key, Decimal("0")) + Decimal(str(r["comision"]))
    return [
        {"vendedor": k[0], "qna": k[1], "total": v}
        for k, v in sorted(agg.items(), key=lambda x: (-x[1], x[0][0]))
    ]


def build_commission_excel(db: Session, **filters) -> Path:
    COMMISSION_EXPORTS_DIR.mkdir(parents=True, exist_ok=True)
    rows = _rows_for_export(db, **filters)
    summary = _summary_by_seller_qna(rows)

    wb = Workbook()
    ws = wb.active
    ws.title = "Comisiones"
    headers = [
        "Folio",
        "Cliente",
        "Vendedor",
        "QNA",
        "Semana",
        "Monto venta",
        "%",
        "Comisión",
        "Estado",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append(
            [
                r["folio"],
                r["cliente"],
                r["vendedor"],
                r["qna"],
                r["semana"],
                float(r["venta"]),
                float(r["pct"]),
                float(r["comision"]),
                r["estado"],
            ]
        )

    ws2 = wb.create_sheet("Resumen vendedor QNA")
    ws2.append(["Vendedor", "QNA", "Total comisión"])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
    for s in summary:
        ws2.append([s["vendedor"], s["qna"], float(s["total"])])

    ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
    path = COMMISSION_EXPORTS_DIR / f"comisiones_{ts}.xlsx"
    wb.save(path)
    return path


def build_commission_pdf(db: Session, **filters) -> bytes:
    rows = _rows_for_export(db, **filters)
    summary = _summary_by_seller_qna(rows)

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Resumen de comisiones — Sistema Gaman", styles["Title"]),
        Spacer(1, 12),
    ]

    total = sum(
        Decimal(str(r["comision"]))
        for r in rows
        if r["estado"] != _STATUS_LABELS[PAYMENT_STATUS_CANCELLED]
    )
    story.append(Paragraph(f"Total comisiones (no canceladas): ${total:,.2f}", styles["Normal"]))
    story.append(Spacer(1, 16))

    if summary:
        sum_data = [["Vendedor", "QNA", "Total"]] + [
            [s["vendedor"], s["qna"], f"${float(s['total']):,.2f}"] for s in summary[:40]
        ]
        t_sum = Table(sum_data, colWidths=[2.5 * inch, 1.2 * inch, 1.5 * inch])
        t_sum.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e293b")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ]
            )
        )
        story.append(t_sum)
        story.append(Spacer(1, 20))

    detail_data = [
        ["Folio", "Cliente", "Vendedor", "QNA", "Venta", "%", "Comisión", "Estado"]
    ]
    for r in rows[:80]:
        detail_data.append(
            [
                str(r["folio"])[:12],
                str(r["cliente"])[:18],
                str(r["vendedor"])[:14],
                str(r["qna"])[:8],
                f"${float(r['venta']):,.0f}",
                f"{float(r['pct']):.0f}%",
                f"${float(r['comision']):,.2f}",
                r["estado"],
            ]
        )
    t_det = Table(
        detail_data,
        colWidths=[0.7 * inch, 1.4 * inch, 1.2 * inch, 0.6 * inch, 0.8 * inch, 0.4 * inch, 0.8 * inch, 0.7 * inch],
    )
    t_det.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#334155")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
            ]
        )
    )
    story.append(t_det)
    doc.build(story)
    return buffer.getvalue()
