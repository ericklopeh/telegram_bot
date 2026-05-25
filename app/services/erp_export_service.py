"""Exportaciones Excel ERP (solo storage/excel_exports/)."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.core.paths import EXCEL_EXPORTS_DIR
from app.models.contract import Contract
from app.models.erp_payment import ErpPayment
from app.services.contract_financial_service import ContractFinancialService
from app.services.erp_reconciliation_service import ErpReconciliationService
from app.services.excel_path_guard import assert_writable_excel_path


class ErpExportService:
    def __init__(self) -> None:
        self.financial = ContractFinancialService()

    def _export_path(self, prefix: str) -> Path:
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        dest = EXCEL_EXPORTS_DIR / "erp" / f"{prefix}_{ts}.xlsx"
        dest.parent.mkdir(parents=True, exist_ok=True)
        assert_writable_excel_path(dest)
        return dest

    def export_executive_summary(self, db: Session) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen ejecutivo"
        ws.append(["Contrato", "Cliente", "Folio", "Venta", "Pagado", "Saldo", "Refin", "Recovery %", "Estado"])
        for cell in ws[1]:
            cell.font = Font(bold=True)

        contracts = db.scalars(
            select(Contract).options(joinedload(Contract.customer)).order_by(Contract.id.desc())
        ).unique().all()
        for c in contracts:
            snap = self.financial.compute_snapshot(db, c)
            ws.append(
                [
                    c.id,
                    c.customer.name if c.customer else c.customer_id,
                    c.folio or "",
                    float(snap.total_sale),
                    float(snap.total_paid),
                    float(snap.saldo_final),
                    float(snap.refinanciado),
                    float(snap.recovery_pct),
                    c.status,
                ]
            )
        path = self._export_path("erp_ejecutivo")
        wb.save(path)
        return path

    def export_cobranza(self, db: Session) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Cobranza"
        ws.append(["Contrato", "Cliente", "Atraso", "Cuotas vencidas", "Saldo"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for c in db.scalars(select(Contract).options(joinedload(Contract.customer))).unique().all():
            snap = self.financial.compute_snapshot(db, c)
            if snap.atraso <= 0:
                continue
            ws.append(
                [
                    c.id,
                    c.customer.name if c.customer else "",
                    float(snap.atraso),
                    snap.overdue_installments,
                    float(snap.saldo_final),
                ]
            )
        path = self._export_path("erp_cobranza")
        wb.save(path)
        return path

    def export_reconciliation(self, db: Session) -> Path:
        report = ErpReconciliationService().build_report(db)
        wb = Workbook()
        ws = wb.active
        ws.title = "Conciliacion"
        ws.append(["Codigo", "Severidad", "Mensaje", "Entidad", "ID"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for issue in report.issues:
            ws.append([issue.code, issue.severity, issue.message, issue.entity_type, issue.entity_id])
        path = self._export_path("erp_conciliacion")
        wb.save(path)
        return path

    def export_recent_payments(self, db: Session, *, limit: int = 500) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Pagos"
        ws.append(["ID", "Venta", "Monto", "Fecha", "Referencia"])
        payments = list(
            db.scalars(select(ErpPayment).order_by(ErpPayment.id.desc()).limit(limit)).all()
        )
        for p in payments:
            ws.append([p.id, p.sale_id, float(p.amount), str(p.payment_date or ""), p.reference or ""])
        path = self._export_path("erp_pagos")
        wb.save(path)
        return path

    def export_vendedor_summary(self, db: Session) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Vendedores"
        ws.append(["Vendedor", "Ventas activas", "Monto venta", "Pagado", "Recovery %"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        kpis = self.financial.build_dashboard_kpis(db)
        for vendedor, monto in kpis.top_vendedores:
            ws.append([vendedor, "", float(monto), "", ""])
        path = self._export_path("erp_vendedores")
        wb.save(path)
        return path

    def export_recovery(self, db: Session) -> Path:
        wb = Workbook()
        ws = wb.active
        ws.title = "Recovery"
        ws.append(["Contrato", "Cliente", "Venta", "Pagado", "Recovery %", "Saldo"])
        for c in db.scalars(select(Contract).options(joinedload(Contract.customer))).unique().all():
            snap = self.financial.compute_snapshot(db, c)
            ws.append(
                [
                    c.id,
                    c.customer.name if c.customer else "",
                    float(snap.total_sale),
                    float(snap.total_paid),
                    float(snap.recovery_pct),
                    float(snap.saldo_final),
                ]
            )
        path = self._export_path("erp_recovery")
        wb.save(path)
        try:
            from app.services.platform_cohesion_service import safe_cohesion_emit

            safe_cohesion_emit(
                db,
                action="export_generated",
                title="Export ERP recovery",
                entity_type="export",
                source="erp",
                href=str(path),
            )
        except Exception:
            pass
        return path

    def export_refinance(self, db: Session) -> Path:
        from app.models.contract import RefinanceOperation

        wb = Workbook()
        ws = wb.active
        ws.title = "Refinanciamientos"
        ws.append(["ID", "Contrato", "Anterior", "Refinanciado", "Nuevo saldo", "Notas", "Fecha"])
        for op in db.scalars(
            select(RefinanceOperation).order_by(RefinanceOperation.id.desc()).limit(1000)
        ).all():
            ws.append(
                [
                    op.id,
                    op.contract_id,
                    float(op.previous_balance),
                    float(op.refinanced_amount),
                    float(op.new_balance),
                    op.notes or "",
                    str(op.created_at or ""),
                ]
            )
        path = self._export_path("erp_refin")
        wb.save(path)
        return path
