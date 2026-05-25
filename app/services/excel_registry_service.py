"""Registro de ventas en copias de Excel maestros (P19) — no modifica originales."""

from __future__ import annotations

import logging
import shutil
import unicodedata
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from app.core.paths import (
    CONTRATOS_MASTER_FILENAME,
    CONTRATOS_MASTER_PATH,
    EXCEL_EXPORTS_DAILY_DIR,
    EXCEL_EXPORTS_OPS_DIR,
    VENTAS_MASTER_FILENAME,
    VENTAS_MASTER_PATH,
)
from app.models.sale_capture import SaleCapture
from app.services.excel_path_guard import assert_writable_excel_path, is_under_excel_masters

log = logging.getLogger(__name__)

_VENTAS_SHEET = "CAPTURA_2026"
_VENTAS_HEADER_ROW = 5
_VENTAS_DATA_START = 6
_VENTAS_MAX_COL = 27

_CONTRATOS_HEADER_MARKERS = ("ESTATUS", "NOMBRE")
_CONTRATOS_SKIP = frozenset(
    {"DASHBOARD", "README", "VENDEDORES", "COMISIONES", "CODIGOS_DINERO", "RESUMEN", "OFICINA"}
)

# Columnas 1-based en CAPTURA_2026
_COL_VENTAS = {
    "fecha": 1,
    "folio": 2,
    "vendedor": 3,
    "tipo_venta": 4,
    "seccion": 5,
    "cliente": 6,
    "codigo": 7,
    "producto": 8,
    "plazo": 9,
    "costo": 10,
    "venta_real": 14,
    "tipo_cotizacion": 17,
    "recuperacion": 19,
    "semana": 20,
    "rfc": 22,
    "observaciones": 26,
}


class ExcelRegistryError(Exception):
    """Error controlado al escribir copias Excel (P20-G)."""

    def __init__(self, message: str, *, code: str | None = None) -> None:
        super().__init__(message)
        self.code = code


def _normalize_sheet_key(name: str) -> str:
    text = unicodedata.normalize("NFKD", name.strip().upper())
    return "".join(ch for ch in text if not unicodedata.combining(ch))


def _resolve_contratos_sheet(wb, vendedor: str) -> str | None:
    target = _normalize_sheet_key(vendedor)
    for name in wb.sheetnames:
        if name in _CONTRATOS_SKIP:
            continue
        if _normalize_sheet_key(name) == target:
            return name
    for name in wb.sheetnames:
        if name in _CONTRATOS_SKIP:
            continue
        if target in _normalize_sheet_key(name):
            return name
    return None


def _find_contratos_header_row(ws, max_scan: int = 25) -> int | None:
    for row_idx in range(1, max_scan + 1):
        labels = []
        for cell in ws[row_idx]:
            if cell.value:
                labels.append(str(cell.value).strip().upper())
        joined = " ".join(labels)
        if "ESTATUS" in joined and ("CLIENTE" in joined or "NOMBRE" in joined):
            return row_idx
    return None


def _next_empty_row_ventas(ws, start: int = _VENTAS_DATA_START) -> int:
    row = start
    while row < start + 50000:
        folio = ws.cell(row, _COL_VENTAS["folio"]).value
        cliente = ws.cell(row, _COL_VENTAS["cliente"]).value
        if folio in (None, "") and cliente in (None, ""):
            return row
        row += 1
    raise ExcelRegistryError("No hay filas vacías disponibles en CAPTURA_2026.")


def _next_empty_row_contratos(ws, header_row: int) -> int:
    row = header_row + 1
    while row < header_row + 5000:
        cliente = ws.cell(row, 3).value
        codigo = ws.cell(row, 6).value
        if cliente in (None, "") and codigo in (None, ""):
            return row
        row += 1
    raise ExcelRegistryError("No hay filas vacías en la hoja de contratos del vendedor.")


def _copy_row_formulas_and_styles(ws, src_row: int, dst_row: int, max_col: int) -> None:
    """Copia estilos y fórmulas de la fila anterior a la nueva, ajustando referencias."""
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)
            dst.alignment = copy(src.alignment)
        if isinstance(src.value, str) and src.value.startswith("="):
            dst.value = Translator(src.value, origin=src.coordinate).translate_formula(
                dst.coordinate
            )


def _write_ventas_values(ws, row: int, sale: SaleCapture, semana: str) -> None:
    ws.cell(row, _COL_VENTAS["fecha"]).value = sale.sale_date
    try:
        ws.cell(row, _COL_VENTAS["folio"]).value = int(sale.folio)
    except ValueError:
        ws.cell(row, _COL_VENTAS["folio"]).value = sale.folio
    ws.cell(row, _COL_VENTAS["vendedor"]).value = sale.vendedor
    ws.cell(row, _COL_VENTAS["tipo_venta"]).value = sale.tipo_venta or "PARALELA"
    ws.cell(row, _COL_VENTAS["seccion"]).value = sale.seccion
    ws.cell(row, _COL_VENTAS["cliente"]).value = sale.cliente
    ws.cell(row, _COL_VENTAS["codigo"]).value = sale.codigo
    ws.cell(row, _COL_VENTAS["producto"]).value = sale.producto
    ws.cell(row, _COL_VENTAS["plazo"]).value = sale.plazo
    if sale.costo is not None:
        ws.cell(row, _COL_VENTAS["costo"]).value = float(sale.costo)
    if sale.venta is not None:
        ws.cell(row, _COL_VENTAS["venta_real"]).value = float(sale.venta)
    elif sale.total_venta is not None:
        ws.cell(row, _COL_VENTAS["venta_real"]).value = float(sale.total_venta)
    ws.cell(row, _COL_VENTAS["tipo_cotizacion"]).value = sale.tipo_cotizacion
    if sale.recuperacion is not None:
        ws.cell(row, _COL_VENTAS["recuperacion"]).value = float(sale.recuperacion)
    ws.cell(row, _COL_VENTAS["semana"]).value = semana
    ws.cell(row, _COL_VENTAS["rfc"]).value = (sale.rfc or "").upper()
    obs_parts = []
    if sale.qna:
        obs_parts.append(f"QNA:{sale.qna}")
    if sale.venta_refinanciamiento is not None:
        obs_parts.append(f"Refi:{sale.venta_refinanciamiento}")
    if sale.venta_sin_agregado is not None:
        obs_parts.append(f"Sin agregado:{sale.venta_sin_agregado}")
    if sale.precio_comision is not None:
        obs_parts.append(f"P.Comisión:{sale.precio_comision}")
    if sale.observaciones:
        obs_parts.append(sale.observaciones)
    if obs_parts:
        ws.cell(row, _COL_VENTAS["observaciones"]).value = " | ".join(obs_parts)


def _write_contratos_values(ws, row: int, sale: SaleCapture) -> None:
    ws.cell(row, 3).value = sale.cliente
    ws.cell(row, 4).value = sale.qna
    ws.cell(row, 5).value = sale.plazo
    ws.cell(row, 6).value = sale.codigo
    ws.cell(row, 7).value = sale.producto
    venta_val = sale.total_venta or sale.venta
    if venta_val is not None:
        ws.cell(row, 8).value = float(venta_val)
    if sale.precio_comision is not None:
        ws.cell(row, 12).value = float(sale.precio_comision)
    ws.cell(row, 13).value = sale.tipo_cotizacion or sale.tipo_venta


class ExcelRegistryService:
    """Copia masters y agrega filas de venta en exportaciones."""

    def __init__(self) -> None:
        EXCEL_EXPORTS_DAILY_DIR.mkdir(parents=True, exist_ok=True)
        EXCEL_EXPORTS_OPS_DIR.mkdir(parents=True, exist_ok=True)

    def daily_dir_for(self, day: date | None = None) -> Path:
        d = day or date.today()
        path = EXCEL_EXPORTS_DAILY_DIR / d.isoformat()
        path.mkdir(parents=True, exist_ok=True)
        return path

    def operation_dir_for(self, sale_id: int) -> Path:
        path = EXCEL_EXPORTS_OPS_DIR / str(sale_id)
        path.mkdir(parents=True, exist_ok=True)
        return path

    def ensure_copy(self, master: Path, dest: Path) -> Path:
        if is_under_excel_masters(dest):
            raise ExcelRegistryError(
                "No se puede escribir sobre el archivo maestro original.",
                code="master_write_forbidden",
            )
        if not master.is_file():
            raise ExcelRegistryError(
                f"Plantilla maestra no encontrada: {master.name}",
                code="template_missing",
            )
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.is_file():
            try:
                shutil.copy2(master, dest)
            except PermissionError as exc:
                raise ExcelRegistryError(
                    "Sin permisos para crear la copia de exportación.",
                    code="permission_denied",
                ) from exc
            except OSError as exc:
                raise ExcelRegistryError(
                    f"No se pudo crear la copia de exportación: {exc}",
                    code="copy_failed",
                ) from exc
            log.info("Copia Excel creada", extra={"dest": str(dest)})
        return dest

    def prepare_operation_workbooks(self, sale: SaleCapture) -> tuple[Path, Path]:
        """Copia fresca por operación (desde master original)."""
        op_dir = self.operation_dir_for(sale.id)
        ventas_dest = op_dir / VENTAS_MASTER_FILENAME
        contratos_dest = op_dir / CONTRATOS_MASTER_FILENAME
        self.ensure_copy(VENTAS_MASTER_PATH, ventas_dest)
        self.ensure_copy(CONTRATOS_MASTER_PATH, contratos_dest)
        return ventas_dest, contratos_dest

    def prepare_daily_workbooks(self) -> tuple[Path, Path]:
        daily = self.daily_dir_for()
        ventas_dest = daily / VENTAS_MASTER_FILENAME
        contratos_dest = daily / CONTRATOS_MASTER_FILENAME
        self.ensure_copy(VENTAS_MASTER_PATH, ventas_dest)
        self.ensure_copy(CONTRATOS_MASTER_PATH, contratos_dest)
        return ventas_dest, contratos_dest

    def append_sale_to_workbooks(
        self,
        sale: SaleCapture,
        *,
        ventas_path: Path,
        contratos_path: Path,
        semana: str,
        use_daily_base: bool = False,
    ) -> dict[str, Any]:
        """
        Agrega fila en copias. Si use_daily_base, parte de la copia diaria acumulada.
        """
        if use_daily_base:
            daily_v, daily_c = self.prepare_daily_workbooks()
            if ventas_path != daily_v:
                shutil.copy2(daily_v, ventas_path)
            if contratos_path != daily_c:
                shutil.copy2(daily_c, contratos_path)

        if not ventas_path.parent.is_dir():
            raise ExcelRegistryError(
                f"Ruta de exportación inexistente: {ventas_path.parent}",
                code="path_missing",
            )
        try:
            v_row = self._append_ventas(ventas_path, sale, semana)
            c_row, c_sheet = self._append_contratos(contratos_path, sale)
        except PermissionError as exc:
            raise ExcelRegistryError(
                "No se puede escribir el Excel (archivo abierto o bloqueado). Cierra Excel e intenta de nuevo.",
                code="file_locked",
            ) from exc
        except OSError as exc:
            if getattr(exc, "errno", None) == 13:
                raise ExcelRegistryError(
                    "Sin permisos para escribir el archivo Excel de exportación.",
                    code="permission_denied",
                ) from exc
            raise ExcelRegistryError(
                f"Error de sistema al escribir Excel: {exc}",
                code="io_error",
            ) from exc

        if use_daily_base:
            shutil.copy2(ventas_path, daily_v)
            shutil.copy2(contratos_path, daily_c)

        return {
            "ventas_path": str(ventas_path),
            "contratos_path": str(contratos_path),
            "ventas_row": v_row,
            "contratos_row": c_row,
            "contratos_sheet": c_sheet,
        }

    def _append_ventas(self, path: Path, sale: SaleCapture, semana: str) -> int:
        if not path.is_file():
            raise ExcelRegistryError(
                f"Archivo de ventas no encontrado: {path}",
                code="path_missing",
            )
        try:
            wb = load_workbook(path, data_only=False)
        except PermissionError as exc:
            raise ExcelRegistryError(
                "El Excel de ventas está abierto o bloqueado. Ciérralo e intenta de nuevo.",
                code="file_locked",
            ) from exc
        except OSError as exc:
            raise ExcelRegistryError(
                f"No se puede abrir el Excel de ventas: {exc}",
                code="io_error",
            ) from exc
        if _VENTAS_SHEET not in wb.sheetnames:
            wb.close()
            raise ExcelRegistryError(
                f"Hoja obligatoria «{_VENTAS_SHEET}» no encontrada en ventas.",
                code="sheet_missing",
            )
        ws = wb[_VENTAS_SHEET]
        if not sale.folio or not sale.cliente:
            wb.close()
            raise ExcelRegistryError(
                "Faltan columnas obligatorias (folio o cliente) para ventas.",
                code="required_column",
            )
        new_row = _next_empty_row_ventas(ws)
        template_row = max(new_row - 1, _VENTAS_DATA_START)
        if new_row > _VENTAS_DATA_START:
            _copy_row_formulas_and_styles(ws, template_row, new_row, _VENTAS_MAX_COL)
        _write_ventas_values(ws, new_row, sale, semana)
        try:
            assert_writable_excel_path(path)
            wb.save(path)
        except PermissionError as exc:
            wb.close()
            raise ExcelRegistryError(
                "No se pudo guardar ventas (archivo abierto). Cierra Excel e intenta de nuevo.",
                code="save_failed",
            ) from exc
        except OSError as exc:
            wb.close()
            raise ExcelRegistryError(
                f"Error al guardar workbook de ventas: {exc}",
                code="save_failed",
            ) from exc
        wb.close()
        return new_row

    def _append_contratos(self, path: Path, sale: SaleCapture) -> tuple[int, str]:
        if not path.is_file():
            raise ExcelRegistryError(
                f"Archivo de contratos no encontrado: {path}",
                code="path_missing",
            )
        try:
            wb = load_workbook(path, data_only=False)
        except PermissionError as exc:
            raise ExcelRegistryError(
                "El Excel de contratos está abierto o bloqueado. Ciérralo e intenta de nuevo.",
                code="file_locked",
            ) from exc
        sheet_name = _resolve_contratos_sheet(wb, sale.vendedor)
        if not sheet_name:
            wb.close()
            raise ExcelRegistryError(
                f"No existe hoja de contratos para el vendedor «{sale.vendedor}».",
                code="sheet_missing",
            )
        ws = wb[sheet_name]
        header_row = _find_contratos_header_row(ws)
        if not header_row:
            wb.close()
            raise ExcelRegistryError(
                f"Encabezados obligatorios no detectados en hoja «{sheet_name}».",
                code="required_column",
            )
        new_row = _next_empty_row_contratos(ws, header_row)
        template_row = max(new_row - 1, header_row + 1)
        if new_row > header_row + 1:
            _copy_row_formulas_and_styles(ws, template_row, new_row, 14)
        ws.cell(new_row, 2).value = "N"
        _write_contratos_values(ws, new_row, sale)
        try:
            assert_writable_excel_path(path)
            wb.save(path)
        except PermissionError as exc:
            wb.close()
            raise ExcelRegistryError(
                "No se pudo guardar contratos (archivo abierto). Cierra Excel e intenta de nuevo.",
                code="save_failed",
            ) from exc
        except OSError as exc:
            wb.close()
            raise ExcelRegistryError(
                f"Error al guardar workbook de contratos: {exc}",
                code="save_failed",
            ) from exc
        wb.close()
        return new_row, sheet_name

    def folio_exists_in_ventas_copy(self, path: Path, folio: str) -> bool:
        if not path.is_file():
            return False
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb[_VENTAS_SHEET]
            needle = str(folio).strip()
            for row in ws.iter_rows(
                min_row=_VENTAS_DATA_START, max_col=2, values_only=True
            ):
                if row and len(row) > 1 and str(row[1]).strip() == needle:
                    return True
            return False
        finally:
            wb.close()

    def suggest_next_folio_from_copy(self, path: Path) -> int:
        if not path.is_file():
            return 0
        wb = load_workbook(path, read_only=True, data_only=True)
        best = 0
        try:
            ws = wb[_VENTAS_SHEET]
            for row in ws.iter_rows(
                min_row=_VENTAS_DATA_START, min_col=2, max_col=2, values_only=True
            ):
                val = row[0] if row else None
                if val is None:
                    continue
                try:
                    best = max(best, int(float(val)))
                except (TypeError, ValueError):
                    continue
        finally:
            wb.close()
        return best
