"""Exportaciones Excel BI (solo storage/excel_exports/bi/)."""

from __future__ import annotations

import logging
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy.orm import Session

from app.core.paths import EXCEL_EXPORTS_DIR
from app.services.bi_dashboard_service import BiDashboardService
from app.services.bi_filters import BiFilters
from app.services.case_event_service import BI_REPORT_EXPORTED
from app.services.excel_path_guard import assert_writable_excel_path

log = logging.getLogger(__name__)


class BiExportService:
    def __init__(self) -> None:
        self.dashboard = BiDashboardService()

    def _export_path(self, name: str) -> Path:
        ts = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        dest = EXCEL_EXPORTS_DIR / "bi" / f"{name}_{ts}.xlsx"
        dest.parent.mkdir(parents=True, exist_ok=True)
        assert_writable_excel_path(dest)
        return dest

    def _log_export(self, export_type: str, path: Path) -> None:
        log.info("BI export %s -> %s", export_type, path)
        log.info("event=%s path=%s", BI_REPORT_EXPORTED, path.name)

    def export_dashboard_resumen(self, db: Session, filters: BiFilters) -> Path:
        payload = self.dashboard.build_dashboard(db, filters)
        wb = Workbook()
        ws = wb.active
        ws.title = "Resumen BI"
        ws.append(["KPI", "Valor", "Detalle", "Enlace"])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for k in payload.kpis:
            ws.append([k.label, k.value, k.hint, k.drill_url])
        ws2 = wb.create_sheet("Alertas")
        ws2.append(["Código", "Severidad", "Título", "Mensaje", "Cantidad"])
        for a in payload.alerts:
            ws2.append([a.code, a.severity, a.title, a.message, a.count])
        path = self._export_path("dashboard_resumen")
        wb.save(path)
        self._log_export("dashboard_resumen", path)
        return path

    def export_ventas_vendedor(self, db: Session, filters: BiFilters) -> Path:
        payload = self.dashboard.build_dashboard(db, filters)
        chart = payload.charts.get("ventas_vendedor", {})
        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas vendedor"
        ws.append(["Vendedor", "Monto"])
        for lbl, val in zip(chart.get("labels", []), chart.get("values", [])):
            ws.append([lbl, val])
        path = self._export_path("ventas_vendedor")
        wb.save(path)
        self._log_export("ventas_vendedor", path)
        return path

    def export_recovery(self, db: Session, filters: BiFilters) -> Path:
        payload = self.dashboard.build_dashboard(db, filters)
        rows = payload.rankings.get("top_recovery", [])
        wb = Workbook()
        ws = wb.active
        ws.title = "Recovery"
        ws.append(["Vendedor", "Recovery %", "Detalle"])
        for r in rows:
            ws.append([r.label, r.metric, r.detail])
        path = self._export_path("recovery")
        wb.save(path)
        self._log_export("recovery", path)
        return path

    def export_pendientes(self, db: Session, filters: BiFilters) -> Path:
        payload = self.dashboard.build_dashboard(db, filters)
        wb = Workbook()
        ws = wb.active
        ws.title = "Pendientes"
        ws.append(["Vendedor", "Saldo pendiente", "Detalle"])
        for r in payload.rankings.get("mas_pendientes", []):
            ws.append([r.label, r.metric, r.detail])
        ws2 = wb.create_sheet("Alertas")
        ws2.append(["Código", "Severidad", "Título", "Cantidad", "Mensaje"])
        for a in payload.alerts:
            ws2.append([a.code, a.severity, a.title, a.count, a.message])
        path = self._export_path("pendientes")
        wb.save(path)
        self._log_export("pendientes", path)
        return path
