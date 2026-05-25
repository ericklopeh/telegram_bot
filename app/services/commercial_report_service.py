"""Lectura de masters Excel (ventas/contratos) y autorizaciones BD para control comercial (P18)."""

from __future__ import annotations

import logging
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from app.core.paths import CONTRATOS_MASTER_PATH, VENTAS_MASTER_PATH
from app.web.services.approved_authorizations import list_approved_authorizations

log = logging.getLogger(__name__)

DEFAULT_ROW_LIMIT = 500

_VENTAS_SKIP_SHEETS = frozenset(
    {
        "README",
        "CONFIG",
        "VENDEDORES",
        "COMISIONES",
        "CODIGOS_DINERO",
        "RESUMEN_SEMANAL",
        "PRESTAMOS_SEMANA",
        "PRESTAMOS_VENDEDOR_SEMANA",
        "RESUMEN_VENDEDOR",
        "RESUMEN_RECUPERACION",
        "DASHBOARD",
        "FACTORES",
    }
)

_CONTRATOS_SKIP_SHEETS = frozenset(
    {
        "DASHBOARD",
        "README",
        "VENDEDORES",
        "COMISIONES",
        "CODIGOS_DINERO",
        "RESUMEN",
        "OFICINA",
    }
)

_VENTAS_PRIMARY_SHEET = "CAPTURA_2026"
_VENTAS_HEADER_MARKERS = ("FOLIO", "VENDEDOR", "CLIENTE")

_CONTRATOS_HEADER_MARKERS = ("ESTATUS", "NOMBRE")

_VENTAS_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "fecha": ("FECHA",),
    "folio": ("FOLIO",),
    "vendedor": ("VENDEDOR",),
    "seccion": ("SECCION", "SECCIÓN"),
    "qna": ("QNA",),
    "cliente": ("CLIENTE",),
    "rfc": ("RFC",),
    "producto": ("ARTICULO", "ARTÍCULO", "CODIGO PRODUCTO", "CÓDIGO PRODUCTO"),
    "tipo_venta": ("TIPO VENTA",),
    "plazo": ("PLAZO",),
    "costo": ("COSTO PRODUCTO",),
    "venta": ("VENTA REAL", "VENTA COMISION", "VENTA COMISIÓN"),
    "comision": ("COMISION", "COMISIÓN"),
    "recuperacion": ("RECUPERACION", "RECUPERACIÓN"),
    "semana": ("SEMANA", "SEM_KEY"),
}

_CONTRATOS_FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "estatus": ("ESTATUS",),
    "cliente": ("NOMBRE DEL CLIENTE", "CLIENTE"),
    "qna": ("QNA",),
    "plazo": ("PLAZO",),
    "folio": ("FOLIO",),
    "contrato": ("CONTRATO", "NO CONTRATO", "NO. CONTRATO"),
    "codigo": ("CODIGO", "CÓDIGO"),
    "articulo": ("ARTICULO", "ARTÍCULO"),
    "monto": ("VENTA REAL", "MONTO", "IMPORTE"),
    "comision": ("PAGO COMISION", "PAGO COMISIÓN", "BASE COMISION"),
    "saldo": ("SALDO", "SALDO REF", "SALDO PENDIENTE"),
    "tipo_venta": ("TIPO VENTA",),
}


@dataclass
class ExcelSheetDiagnostic:
    sheet_name: str
    header_row: int | None = None
    columns_detected: dict[str, str] = field(default_factory=dict)
    columns_missing: list[str] = field(default_factory=list)
    rows_read: int = 0
    note: str | None = None


@dataclass
class ExcelReadResult:
    rows: list[dict[str, Any]]
    diagnostics: list[ExcelSheetDiagnostic]
    error: str | None = None
    truncated: bool = False
    source_path: str | None = None


def _normalize_label(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text)


def _format_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        if value == int(value):
            return int(value)
        return round(value, 2)
    return value


def _format_money(value: Any) -> str:
    if value in (None, ""):
        return "—"
    try:
        num = float(value)
        return f"${num:,.2f}"
    except (TypeError, ValueError):
        return str(value)


def _open_workbook(path: Path):
    if not path.is_file():
        raise FileNotFoundError(f"No se encontró el archivo: {path}")
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except PermissionError as exc:
        raise PermissionError(
            f"El archivo está abierto o bloqueado. Cierra Excel e intenta de nuevo: {path.name}"
        ) from exc
    except OSError as exc:
        if getattr(exc, "errno", None) in (13, 32):
            raise PermissionError(
                f"No se puede leer el archivo (¿abierto en Excel?): {path.name}"
            ) from exc
        raise


def _find_header_row(
    ws,
    *,
    markers: tuple[str, ...],
    max_scan: int = 40,
) -> tuple[int | None, list[str]]:
    for row_idx, row in enumerate(
        ws.iter_rows(min_row=1, max_row=max_scan, values_only=True),
        start=1,
    ):
        labels = [_normalize_label(c) for c in (row or ())]
        joined = " | ".join(labels)
        if all(marker in joined for marker in markers):
            return row_idx, labels
    return None, []


def _map_columns(labels: list[str], aliases: dict[str, tuple[str, ...]]) -> dict[str, int | None]:
    index_by_label = {_normalize_label(lbl): idx for idx, lbl in enumerate(labels) if lbl}
    mapping: dict[str, int | None] = {}
    for field, candidates in aliases.items():
        idx = None
        for cand in candidates:
            norm = _normalize_label(cand)
            if norm in index_by_label:
                idx = index_by_label[norm]
                break
        mapping[field] = idx
    return mapping


def _diagnostic_from_map(
    sheet_name: str,
    header_row: int | None,
    col_map: dict[str, int | None],
    aliases: dict[str, tuple[str, ...]],
    labels: list[str],
    *,
    rows_read: int = 0,
    note: str | None = None,
) -> ExcelSheetDiagnostic:
    detected = {}
    missing = []
    for field in aliases:
        idx = col_map.get(field)
        if idx is not None and idx < len(labels):
            detected[field] = labels[idx] or f"col_{idx + 1}"
        else:
            missing.append(field)
    return ExcelSheetDiagnostic(
        sheet_name=sheet_name,
        header_row=header_row,
        columns_detected=detected,
        columns_missing=missing,
        rows_read=rows_read,
        note=note,
    )


def _get_cell(row: tuple, idx: int | None) -> Any:
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _matches_filters(
    row: dict[str, Any],
    *,
    vendedor: str | None,
    qna: str | None,
    semana: str | None,
    cliente: str | None,
    rfc: str | None,
    tipo_venta: str | None,
    folio: str | None,
) -> bool:
    if vendedor and vendedor.strip():
        if vendedor.strip().upper() not in str(row.get("vendedor") or "").upper():
            return False
    if qna and qna.strip():
        q_val = str(row.get("qna") or row.get("plazo") or "")
        if qna.strip() not in q_val:
            return False
    if semana and semana.strip():
        sem = str(row.get("semana") or "")
        if semana.strip().upper() not in sem.upper():
            return False
    if cliente and cliente.strip():
        if cliente.strip().upper() not in str(row.get("cliente") or "").upper():
            return False
    if rfc and rfc.strip():
        if rfc.strip().upper() not in str(row.get("rfc") or "").upper():
            return False
    if tipo_venta and tipo_venta.strip():
        if tipo_venta.strip().upper() not in str(row.get("tipo_venta") or "").upper():
            return False
    if folio and folio.strip():
        needle = folio.strip().upper()
        hay = " ".join(
            str(row.get(k) or "")
            for k in ("folio", "contrato", "codigo", "cliente")
        ).upper()
        if needle not in hay:
            return False
    return True


def _list_ventas_sheet_names(path: Path) -> list[str]:
    wb = _open_workbook(path)
    try:
        return [
            n
            for n in wb.sheetnames
            if n not in _VENTAS_SKIP_SHEETS and not n.startswith("SEM ")
        ] + [n for n in wb.sheetnames if n.startswith("SEM ")]
    finally:
        wb.close()


def _list_contratos_vendor_sheets(path: Path) -> list[str]:
    wb = _open_workbook(path)
    try:
        return [n for n in wb.sheetnames if n not in _CONTRATOS_SKIP_SHEETS]
    finally:
        wb.close()


def _parse_ventas_row(row: tuple, col_map: dict[str, int | None]) -> dict[str, Any] | None:
    folio = _get_cell(row, col_map.get("folio"))
    cliente = _get_cell(row, col_map.get("cliente"))
    if folio in (None, "") and cliente in (None, ""):
        return None
    qna_val = _get_cell(row, col_map.get("qna"))
    if qna_val in (None, ""):
        qna_val = _get_cell(row, col_map.get("plazo"))
    producto = _get_cell(row, col_map.get("producto"))
    costo = _get_cell(row, col_map.get("costo"))
    venta = _get_cell(row, col_map.get("venta"))
    comision = _get_cell(row, col_map.get("comision"))
    recuperacion = _get_cell(row, col_map.get("recuperacion"))
    return {
        "fecha": _format_cell(_get_cell(row, col_map.get("fecha"))),
        "folio": _format_cell(folio) if folio not in (None, "") else "—",
        "vendedor": _get_cell(row, col_map.get("vendedor")) or "—",
        "seccion": _get_cell(row, col_map.get("seccion")) or "—",
        "qna": _format_cell(qna_val) if qna_val not in (None, "") else "—",
        "cliente": cliente or "—",
        "rfc": _get_cell(row, col_map.get("rfc")) or "—",
        "producto": producto or "—",
        "tipo_venta": _get_cell(row, col_map.get("tipo_venta")) or "—",
        "plazo": _format_cell(_get_cell(row, col_map.get("plazo"))),
        "costo": _format_money(costo),
        "venta": _format_money(venta),
        "comision": _format_money(comision),
        "recuperacion": _format_money(recuperacion),
        "semana": _get_cell(row, col_map.get("semana")) or "—",
        "_sort_folio": folio,
    }


def _parse_contratos_row(
    row: tuple,
    col_map: dict[str, int | None],
    *,
    vendedor_hoja: str,
) -> dict[str, Any] | None:
    cliente = _get_cell(row, col_map.get("cliente"))
    codigo = _get_cell(row, col_map.get("codigo"))
    articulo = _get_cell(row, col_map.get("articulo"))
    if cliente in (None, "") and codigo in (None, "") and articulo in (None, ""):
        return None
    folio = _get_cell(row, col_map.get("folio")) or codigo
    contrato = _get_cell(row, col_map.get("contrato")) or articulo or codigo
    monto = _get_cell(row, col_map.get("monto"))
    comision = _get_cell(row, col_map.get("comision"))
    saldo = _get_cell(row, col_map.get("saldo"))
    return {
        "vendedor": vendedor_hoja,
        "cliente": cliente or "—",
        "folio": _format_cell(folio) if folio not in (None, "") else "—",
        "contrato": contrato or "—",
        "monto": _format_money(monto),
        "plazo": _format_cell(_get_cell(row, col_map.get("plazo"))),
        "comision": _format_money(comision),
        "qna": _format_cell(_get_cell(row, col_map.get("qna"))),
        "saldo": _format_money(saldo) if saldo not in (None, "") else "—",
        "estatus": _get_cell(row, col_map.get("estatus")) or "—",
        "tipo_venta": _get_cell(row, col_map.get("tipo_venta")) or "—",
    }


def _iter_sheet_data_rows(ws, header_row: int) -> Iterator[tuple]:
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        yield row


def load_ventas_from_path(
    path: Path,
    *,
    sheet_name: str | None = None,
    vendedor: str | None = None,
    qna: str | None = None,
    semana: str | None = None,
    cliente: str | None = None,
    rfc: str | None = None,
    tipo_venta: str | None = None,
    limit: int = DEFAULT_ROW_LIMIT,
) -> ExcelReadResult:
    """Lee ventas desde cualquier ruta (imports o copia); no escribe en masters."""
    from app.services.excel_path_guard import assert_not_excel_master_path

    assert_not_excel_master_path(path)
    result = ExcelReadResult(rows=[], diagnostics=[], source_path=str(path))
    sheet = sheet_name or (
        semana.strip() if semana and semana.strip().startswith("SEM ") else _VENTAS_PRIMARY_SHEET
    )

    try:
        wb = _open_workbook(path)
    except FileNotFoundError as exc:
        result.error = str(exc)
        return result
    except PermissionError as exc:
        result.error = str(exc)
        return result

    try:
        if sheet not in wb.sheetnames:
            result.error = (
                f"Hoja «{sheet}» no encontrada en {path.name}. "
                f"Hojas disponibles (muestra): {', '.join(wb.sheetnames[:8])}…"
            )
            return result

        ws = wb[sheet]
        header_row, labels = _find_header_row(ws, markers=_VENTAS_HEADER_MARKERS)
        if not header_row:
            result.error = f"No se detectó fila de encabezados en hoja «{sheet}»."
            result.diagnostics.append(
                ExcelSheetDiagnostic(
                    sheet_name=sheet,
                    note="Se esperaban columnas FOLIO, VENDEDOR y CLIENTE.",
                )
            )
            return result

        col_map = _map_columns(labels, _VENTAS_FIELD_ALIASES)
        diag = _diagnostic_from_map(
            sheet, header_row, col_map, _VENTAS_FIELD_ALIASES, labels
        )
        result.diagnostics.append(diag)

        matched: list[dict[str, Any]] = []
        for row in _iter_sheet_data_rows(ws, header_row):
            parsed = _parse_ventas_row(row, col_map)
            if not parsed:
                continue
            if not _matches_filters(
                parsed,
                vendedor=vendedor,
                qna=qna,
                semana=semana if sheet == _VENTAS_PRIMARY_SHEET else None,
                cliente=cliente,
                rfc=rfc,
                tipo_venta=tipo_venta,
                folio=None,
            ):
                continue
            matched.append(parsed)
            if len(matched) >= limit:
                result.truncated = True
                break

        diag.rows_read = len(matched)
        result.rows = matched
        return result
    finally:
        wb.close()


def load_ventas_2026(
    *,
    vendedor: str | None = None,
    qna: str | None = None,
    semana: str | None = None,
    cliente: str | None = None,
    rfc: str | None = None,
    tipo_venta: str | None = None,
    limit: int = DEFAULT_ROW_LIMIT,
) -> ExcelReadResult:
    return load_ventas_from_path(
        VENTAS_MASTER_PATH,
        vendedor=vendedor,
        qna=qna,
        semana=semana,
        cliente=cliente,
        rfc=rfc,
        tipo_venta=tipo_venta,
        limit=limit,
    )


def load_contratos_from_path(
    path: Path,
    *,
    vendedor: str | None = None,
    qna: str | None = None,
    cliente: str | None = None,
    folio: str | None = None,
    limit: int = DEFAULT_ROW_LIMIT,
) -> ExcelReadResult:
    from app.services.excel_path_guard import assert_not_excel_master_path

    assert_not_excel_master_path(path)
    result = ExcelReadResult(rows=[], diagnostics=[], source_path=str(path))

    try:
        wb = _open_workbook(path)
    except FileNotFoundError as exc:
        result.error = str(exc)
        return result
    except PermissionError as exc:
        result.error = str(exc)
        return result

    try:
        vendor_sheets = [n for n in wb.sheetnames if n not in _CONTRATOS_SKIP_SHEETS]
        if vendedor and vendedor.strip():
            needle = vendedor.strip().upper()
            vendor_sheets = [n for n in vendor_sheets if needle in n.upper()]

        matched: list[dict[str, Any]] = []
        for sheet_name in vendor_sheets:
            if len(matched) >= limit:
                result.truncated = True
                break
            ws = wb[sheet_name]
            header_row, labels = _find_header_row(
                ws,
                markers=_CONTRATOS_HEADER_MARKERS,
            )
            if not header_row:
                result.diagnostics.append(
                    ExcelSheetDiagnostic(
                        sheet_name=sheet_name,
                        note="Sin encabezados de detalle; hoja omitida.",
                    )
                )
                continue

            col_map = _map_columns(labels, _CONTRATOS_FIELD_ALIASES)
            diag = _diagnostic_from_map(
                sheet_name,
                header_row,
                col_map,
                _CONTRATOS_FIELD_ALIASES,
                labels,
            )
            sheet_count = 0
            for row in _iter_sheet_data_rows(ws, header_row):
                parsed = _parse_contratos_row(row, col_map, vendedor_hoja=sheet_name)
                if not parsed:
                    continue
                if not _matches_filters(
                    parsed,
                    vendedor=vendedor,
                    qna=qna,
                    semana=None,
                    cliente=cliente,
                    rfc=None,
                    tipo_venta=None,
                    folio=folio,
                ):
                    continue
                matched.append(parsed)
                sheet_count += 1
                if len(matched) >= limit:
                    result.truncated = True
                    break
            diag.rows_read = sheet_count
            result.diagnostics.append(diag)

        result.rows = matched
        if not vendor_sheets and not result.error:
            result.error = "No hay hojas de vendedor en el archivo de contratos."
        return result
    finally:
        wb.close()


def load_relacion_contratos(
    *,
    vendedor: str | None = None,
    qna: str | None = None,
    cliente: str | None = None,
    folio: str | None = None,
    limit: int = DEFAULT_ROW_LIMIT,
) -> ExcelReadResult:
    return load_contratos_from_path(
        CONTRATOS_MASTER_PATH,
        vendedor=vendedor,
        qna=qna,
        cliente=cliente,
        folio=folio,
        limit=limit,
    )


def load_filter_options_ventas() -> dict[str, Any]:
    """Opciones de filtro a partir del master de ventas (sin cargar todas las filas)."""
    path = VENTAS_MASTER_PATH
    vendedores: set[str] = set()
    semanas: set[str] = set()
    tipos: set[str] = set()
    try:
        wb = _open_workbook(path)
    except (FileNotFoundError, PermissionError) as exc:
        return {"vendedores": [], "semanas": [], "tipos_venta": [], "error": str(exc)}

    try:
        if _VENTAS_PRIMARY_SHEET not in wb.sheetnames:
            return {"vendedores": [], "semanas": [], "tipos_venta": [], "error": "Hoja CAPTURA_2026 no encontrada."}
        ws = wb[_VENTAS_PRIMARY_SHEET]
        header_row, labels = _find_header_row(ws, markers=_VENTAS_HEADER_MARKERS)
        if not header_row:
            return {"vendedores": [], "semanas": [], "tipos_venta": [], "error": "Encabezados no detectados."}
        col_map = _map_columns(labels, _VENTAS_FIELD_ALIASES)
        scanned = 0
        for row in _iter_sheet_data_rows(ws, header_row):
            scanned += 1
            if scanned > 3000:
                break
            v = _get_cell(row, col_map.get("vendedor"))
            if v:
                vendedores.add(str(v).strip())
            s = _get_cell(row, col_map.get("semana"))
            if s:
                semanas.add(str(s).strip())
            t = _get_cell(row, col_map.get("tipo_venta"))
            if t:
                tipos.add(str(t).strip())
        sem_list = sorted(semanas, reverse=True)
        if not sem_list:
            sem_list = [n for n in wb.sheetnames if n.startswith("SEM ")]
        return {
            "vendedores": sorted(vendedores),
            "semanas": sem_list,
            "tipos_venta": sorted(tipos),
            "error": None,
        }
    finally:
        wb.close()


def load_filter_options_contratos() -> dict[str, Any]:
    path = CONTRATOS_MASTER_PATH
    try:
        wb = _open_workbook(path)
    except (FileNotFoundError, PermissionError) as exc:
        return {"vendedores": [], "error": str(exc)}
    try:
        vendors = [n for n in wb.sheetnames if n not in _CONTRATOS_SKIP_SHEETS]
        return {"vendedores": sorted(vendors), "error": None}
    finally:
        wb.close()


def load_generated_authorizations(
    db: Session,
    usuario: dict | None,
    *,
    q: str | None = None,
    vendedor: str | None = None,
    qna: str | None = None,
    fecha_desde: str | None = None,
    fecha_hasta: str | None = None,
    limit: int = DEFAULT_ROW_LIMIT,
) -> list[dict[str, Any]]:
    rows = list_approved_authorizations(
        db,
        usuario,
        q=q,
        vendedor=vendedor,
        qna=qna,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
    )
    return rows[:limit]


def build_commercial_report_payload(
    db: Session,
    usuario: dict | None,
    *,
    tab: str = "ventas",
    ventas_vendedor: str | None = None,
    ventas_qna: str | None = None,
    ventas_semana: str | None = None,
    ventas_cliente: str | None = None,
    ventas_rfc: str | None = None,
    ventas_tipo: str | None = None,
    contratos_vendedor: str | None = None,
    contratos_qna: str | None = None,
    contratos_cliente: str | None = None,
    contratos_folio: str | None = None,
    auth_q: str | None = None,
    auth_vendedor: str | None = None,
    auth_qna: str | None = None,
    auth_desde: str | None = None,
    auth_hasta: str | None = None,
) -> dict[str, Any]:
    from app.web.services.approved_authorizations import list_filter_options

    active = tab if tab in ("ventas", "contratos", "autorizaciones") else "ventas"
    empty_excel = ExcelReadResult(rows=[], diagnostics=[])

    ventas = (
        load_ventas_2026(
            vendedor=ventas_vendedor,
            qna=ventas_qna,
            semana=ventas_semana,
            cliente=ventas_cliente,
            rfc=ventas_rfc,
            tipo_venta=ventas_tipo,
        )
        if active == "ventas"
        else empty_excel
    )
    contratos = (
        load_relacion_contratos(
            vendedor=contratos_vendedor,
            qna=contratos_qna,
            cliente=contratos_cliente,
            folio=contratos_folio,
        )
        if active == "contratos"
        else empty_excel
    )
    autorizaciones = (
        load_generated_authorizations(
            db,
            usuario,
            q=auth_q,
            vendedor=auth_vendedor,
            qna=auth_qna,
            fecha_desde=auth_desde,
            fecha_hasta=auth_hasta,
        )
        if active == "autorizaciones"
        else []
    )

    return {
        "tab": active,
        "ventas": ventas,
        "contratos": contratos,
        "autorizaciones": autorizaciones,
        "opciones_ventas": (
            load_filter_options_ventas()
            if active == "ventas"
            else {"vendedores": [], "semanas": [], "tipos_venta": [], "error": None}
        ),
        "opciones_contratos": (
            load_filter_options_contratos()
            if active == "contratos"
            else {"vendedores": [], "error": None}
        ),
        "opciones_auth": (
            list_filter_options(db, usuario)
            if active == "autorizaciones"
            else {"vendedores": [], "estados": [], "tipos_venta": []}
        ),
        "limit": DEFAULT_ROW_LIMIT,
        "master_ventas_path": str(VENTAS_MASTER_PATH),
        "master_contratos_path": str(CONTRATOS_MASTER_PATH),
    }


def _decimal_cell(value: Any):
    from decimal import Decimal, InvalidOperation

    if value in (None, ""):
        return None
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    if isinstance(value, Decimal):
        return value
    text = str(value).replace("$", "").replace(",", "").strip()
    if not text or text == "—":
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_date_cell(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value in (None, ""):
        return None
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text[:10], fmt).date()
        except ValueError:
            continue
    return None


def read_ventas_rows_raw(
    path: Path,
    *,
    sheet_name: str | None = None,
    limit: int = 5000,
) -> tuple[list[dict[str, Any]], ExcelReadResult]:
    """Filas normalizadas para importación histórica (P32)."""
    from app.services.excel_path_guard import assert_not_excel_master_path

    assert_not_excel_master_path(path)
    display = load_ventas_from_path(path, sheet_name=sheet_name, limit=limit)
    if display.error:
        return [], display

    sheet = sheet_name or _VENTAS_PRIMARY_SHEET
    rows_out: list[dict[str, Any]] = []
    try:
        wb = _open_workbook(path)
        ws = wb[sheet]
        header_row, labels = _find_header_row(ws, markers=_VENTAS_HEADER_MARKERS)
        if not header_row:
            display.error = "Encabezados no detectados."
            return [], display
        col_map = _map_columns(labels, _VENTAS_FIELD_ALIASES)
        row_num = header_row
        for row in _iter_sheet_data_rows(ws, header_row):
            row_num += 1
            folio = _get_cell(row, col_map.get("folio"))
            cliente = _get_cell(row, col_map.get("cliente"))
            if folio in (None, "") and cliente in (None, ""):
                continue
            venta = _decimal_cell(_get_cell(row, col_map.get("venta")))
            if venta is None:
                venta = _decimal_cell(_get_cell(row, col_map.get("costo")))
            rows_out.append(
                {
                    "row_number": row_num,
                    "sheet_name": sheet,
                    "folio": str(folio).strip() if folio not in (None, "") else "",
                    "cliente": str(cliente or "").strip(),
                    "vendedor": str(_get_cell(row, col_map.get("vendedor")) or "").strip(),
                    "seccion": str(_get_cell(row, col_map.get("seccion")) or "").strip() or None,
                    "rfc": str(_get_cell(row, col_map.get("rfc")) or "").strip() or None,
                    "qna": str(_get_cell(row, col_map.get("qna")) or "").strip() or None,
                    "tipo_venta": str(_get_cell(row, col_map.get("tipo_venta")) or "").strip() or None,
                    "semana": str(_get_cell(row, col_map.get("semana")) or "").strip() or None,
                    "sale_date": _parse_date_cell(_get_cell(row, col_map.get("fecha"))),
                    "total_amount": venta if venta is not None else _decimal_cell("0"),
                }
            )
            if len(rows_out) >= limit:
                break
        wb.close()
    except Exception as exc:
        display.error = str(exc)
        return [], display
    return rows_out, display


def read_contratos_rows_raw(path: Path, *, limit: int = 5000) -> tuple[list[dict[str, Any]], ExcelReadResult]:
    from app.services.excel_path_guard import assert_not_excel_master_path

    assert_not_excel_master_path(path)
    display = load_contratos_from_path(path, limit=limit)
    if display.error:
        return [], display
    rows_out: list[dict[str, Any]] = []
    try:
        wb = _open_workbook(path)
        for sheet_name in [n for n in wb.sheetnames if n not in _CONTRATOS_SKIP_SHEETS]:
            ws = wb[sheet_name]
            header_row, labels = _find_header_row(ws, markers=_CONTRATOS_HEADER_MARKERS)
            if not header_row:
                continue
            col_map = _map_columns(labels, _CONTRATOS_FIELD_ALIASES)
            row_num = header_row
            for row in _iter_sheet_data_rows(ws, header_row):
                row_num += 1
                cliente = _get_cell(row, col_map.get("cliente"))
                codigo = _get_cell(row, col_map.get("codigo"))
                if cliente in (None, "") and codigo in (None, ""):
                    continue
                folio = _get_cell(row, col_map.get("folio")) or codigo
                contrato = _get_cell(row, col_map.get("contrato")) or codigo
                monto = _decimal_cell(_get_cell(row, col_map.get("monto")))
                saldo = _decimal_cell(_get_cell(row, col_map.get("saldo")))
                rows_out.append(
                    {
                        "row_number": row_num,
                        "sheet_name": sheet_name,
                        "folio": str(folio or "").strip(),
                        "contract_code": str(contrato or "").strip() or None,
                        "cliente": str(cliente or "").strip(),
                        "vendedor": sheet_name.strip(),
                        "qna": str(_get_cell(row, col_map.get("qna")) or "").strip() or None,
                        "total_amount": monto if monto is not None else (saldo if saldo is not None else _decimal_cell("0")),
                        "saldo": saldo,
                    }
                )
                if len(rows_out) >= limit:
                    break
            if len(rows_out) >= limit:
                break
        wb.close()
    except Exception as exc:
        display.error = str(exc)
        return [], display
    return rows_out, display
